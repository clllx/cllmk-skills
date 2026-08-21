from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from placeholder_spec import (  # noqa: E402
    STANDARD_COMMON,
    STANDARD_SIGNATURE,
    STANDARD_UNVERIFIED,
    STANDARD_VARIABLES,
    build_placeholder,
    decode_placeholder_name,
    encode_placeholder_name,
)

BUILD = SCRIPT_DIR / "build_offer_docx.py"
VALIDATE = SCRIPT_DIR / "validate_placeholders.py"

# 与主文档附录 B 对齐的示例字段清单。ID 均为示例，不对应任何真实租户。
SAMPLE_FIELDS = [
    {"id": 111001, "name": "年假天数", "type": 10, "isVisible": True, "orgId": "example-org"},
    {"id": 111002, "name": "Annual Leave", "type": 10, "isVisible": True, "orgId": "example-org"},
    {"id": 111003, "name": "隐藏字段", "type": 1, "isVisible": False, "orgId": "example-org"},
    {"id": 111004, "name": "审批附件", "type": 11, "isVisible": True, "orgId": "example-org"},
    {"id": 111005, "name": "入职地点", "type": 1, "isVisible": True, "orgId": "example-org"},
    {"id": 111006, "name": "重名字段", "type": 1, "isVisible": True, "orgId": "example-org"},
    {"id": 111007, "name": "重名字段", "type": 6, "isVisible": True, "orgId": "example-org"},
]


def write_fields(tmpdir: Path) -> Path:
    path = tmpdir / "fields.json"
    # 模拟 cllmk 的双层包装，验证 load_fields_json 的下钻逻辑
    path.write_text(
        json.dumps({"code": 0, "data": {"code": 0, "success": True, "data": SAMPLE_FIELDS}}),
        encoding="utf-8",
    )
    return path


def run_validate(*args: str) -> tuple[int, dict]:
    proc = subprocess.run(
        [sys.executable, str(VALIDATE), *args], capture_output=True, text=True
    )
    return proc.returncode, json.loads(proc.stdout)


def find(report: dict, raw: str) -> dict:
    for item in report["placeholders"]:
        if item["raw"] == raw:
            return item
    raise AssertionError(f"{raw!r} not in report")


class PlaceholderSpecTests(unittest.TestCase):
    def test_enumeration_has_37_entries_and_no_duplicates(self) -> None:
        self.assertEqual(len(STANDARD_VARIABLES), 37)
        self.assertEqual(len(set(STANDARD_VARIABLES)), 37)
        self.assertEqual(len(STANDARD_COMMON), 17)
        self.assertEqual(len(STANDARD_SIGNATURE), 17)
        self.assertEqual(len(STANDARD_UNVERIFIED), 3)

    def test_only_spaces_are_encoded(self) -> None:
        self.assertEqual(encode_placeholder_name("Annual Leave"), "Annual%20Leave")
        self.assertEqual(encode_placeholder_name("A  B"), "A%20%20B")
        # 括号、斜杠、连字符、中文标点一律原样——不是 quote/encodeURIComponent
        for name in ("基本薪资（元/月）", "转正后基本薪资(科桥)", "管理职级-1", "薪资/绩效"):
            self.assertEqual(encode_placeholder_name(name), name)

    def test_decode_is_inverse(self) -> None:
        for name in ("Annual Leave", "A  B", "无空格"):
            self.assertEqual(decode_placeholder_name(encode_placeholder_name(name)), name)

    def test_build_placeholder_omits_id_for_standard_variables(self) -> None:
        self.assertEqual(build_placeholder("候选人姓名"), "{候选人姓名}")
        self.assertEqual(build_placeholder("Annual Leave", 111002), "{Annual%20Leave[111002]}")


class BuildDocxTests(unittest.TestCase):
    def test_dry_run_encodes_spaces_inside_placeholders_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            spec = Path(tmp) / "body.md"
            spec.write_text(
                "# 录用通知书\n\n"
                "尊敬的 {候选人姓名}：\n\n"
                "年假：{Annual Leave[111002]}\n"
                "- 已编码：{Annual%20Leave[111002]}\n",
                encoding="utf-8",
            )
            proc = subprocess.run(
                [sys.executable, str(BUILD), "--spec", str(spec), "--dry-run"],
                capture_output=True,
                text=True,
            )
            self.assertEqual(proc.returncode, 0, proc.stderr)
            report = json.loads(proc.stdout)
            self.assertEqual(
                report["placeholders"], ["候选人姓名", "Annual%20Leave[111002]"]
            )
            # 段落里「尊敬的 」的空格不受影响，只编码花括号内部
            self.assertEqual(report["placeholder_count"], 3)

    def test_generated_docx_keeps_each_placeholder_in_one_run(self) -> None:
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx not installed")
        with tempfile.TemporaryDirectory() as tmp:
            spec = Path(tmp) / "body.md"
            out = Path(tmp) / "out.docx"
            spec.write_text("尊敬的 {候选人姓名}，年假 {Annual Leave[111002]}\n", encoding="utf-8")
            subprocess.run(
                [sys.executable, str(BUILD), "--spec", str(spec), "--out", str(out)],
                capture_output=True,
                text=True,
                check=True,
            )
            from docx import Document

            doc = Document(str(out))
            body = [p for p in doc.paragraphs if p.text.strip()]
            self.assertEqual(len(body), 1)
            # 单 run 是服务端能提取占位符的前提
            self.assertEqual(len(body[0].runs), 1)
            self.assertIn("{Annual%20Leave[111002]}", body[0].text)


class ValidateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.fields = write_fields(Path(self.tmp.name))

    def validate(self, placeholders: str, template_type: str = "0") -> tuple[int, dict]:
        return run_validate(
            "--fields-json", str(self.fields),
            "--placeholders", placeholders,
            "--template-type", template_type,
        )

    def test_accepts_standard_and_custom_placeholders(self) -> None:
        code, report = self.validate("候选人姓名,Annual%20Leave[111002],年假天数[111001]")
        self.assertEqual(code, 0)
        self.assertTrue(report["ok"])
        self.assertEqual(report["invalid_count"], 0)
        self.assertEqual(find(report, "候选人姓名")["kind"], "standard")
        self.assertEqual(find(report, "Annual%20Leave[111002]")["kind"], "custom")

    def test_rejects_custom_field_without_id_and_suggests_the_fix(self) -> None:
        code, report = self.validate("年假天数")
        self.assertEqual(code, 1)
        item = find(report, "年假天数")
        self.assertEqual(item["kind"], "missing_field_id")
        self.assertFalse(item["valid"])
        self.assertIn("{年假天数[111001]}", item["note"])

    def test_rejects_unencoded_space(self) -> None:
        code, report = self.validate("Annual Leave[111002]")
        self.assertEqual(code, 1)
        item = find(report, "Annual Leave[111002]")
        self.assertFalse(item["valid"])
        self.assertIn("{Annual%20Leave[111002]}", item["note"])

    def test_rejects_unknown_name_and_says_server_will_not_reject_it(self) -> None:
        code, report = self.validate("并不存在的字段")
        self.assertEqual(code, 1)
        item = find(report, "并不存在的字段")
        self.assertEqual(item["kind"], "unknown")
        self.assertIn("occurrences", item["note"])

    def test_rejects_id_name_mismatch(self) -> None:
        code, report = self.validate("年假天数[111002]")
        self.assertEqual(code, 1)
        item = find(report, "年假天数[111002]")
        self.assertFalse(item["valid"])
        self.assertIn("Annual%20Leave[111002]", item["note"])

    def test_rejects_missing_field_id(self) -> None:
        code, report = self.validate("年假天数[999999]")
        self.assertEqual(code, 1)
        note = find(report, "年假天数[999999]")["note"]
        # 未知 ID 有两种成因，提示必须都覆盖：场景切换（社招/校招字段 ID 不相交）
        # 与跨租户复用。只提后者会把人引去改本来正确的 ID。
        self.assertIn("hireMode", note)
        self.assertIn("跨租户", note)

    def test_rejects_non_placeholder_field_types(self) -> None:
        code, report = self.validate("审批附件[111004]")
        self.assertEqual(code, 1)
        self.assertIn("type=11", find(report, "审批附件[111004]")["note"])

    def test_hidden_field_is_valid_but_flagged(self) -> None:
        code, report = self.validate("隐藏字段[111003]")
        self.assertEqual(code, 0)
        item = find(report, "隐藏字段[111003]")
        self.assertTrue(item["valid"])
        self.assertIn("隐藏", item["note"])

    def test_standard_and_custom_homonyms_coexist(self) -> None:
        # 「入职地点」既是标准变量，也是租户里的自定义字段名
        code, report = self.validate("入职地点,入职地点[111005]")
        self.assertEqual(code, 0)
        self.assertEqual(find(report, "入职地点")["kind"], "standard")
        self.assertEqual(find(report, "入职地点[111005]")["kind"], "custom")

    def test_duplicate_field_names_list_every_candidate_id(self) -> None:
        code, report = self.validate("重名字段")
        self.assertEqual(code, 1)
        note = find(report, "重名字段")["note"]
        self.assertIn("111006", note)
        self.assertIn("111007", note)

    def test_unverified_standard_variable_is_flagged(self) -> None:
        code, report = self.validate("候选人最高学历")
        self.assertEqual(code, 0)
        self.assertIn("未经实测", find(report, "候选人最高学历")["note"])

    def test_signature_variable_warns_when_esign_is_off(self) -> None:
        code, report = self.validate("个人签署区", template_type="0")
        self.assertEqual(code, 0)
        self.assertTrue(any("不会生效" in w for w in report["warnings"]))

    def test_esign_on_without_personal_signature_area_warns(self) -> None:
        code, report = self.validate("候选人姓名", template_type="1")
        self.assertEqual(code, 0)
        self.assertTrue(any("个人签署区" in w for w in report["warnings"]))

    def test_esign_on_with_personal_signature_area_is_clean(self) -> None:
        code, report = self.validate("个人签署区,公章签署区", template_type="1")
        self.assertEqual(code, 0)
        self.assertEqual(report["warnings"], [])


if __name__ == "__main__":
    unittest.main()
