#!/usr/bin/env python3
from __future__ import annotations

"""
校验 Offer 附件模板 docx（或占位符清单）里的占位符是否合法。

**为什么必须单独校验**：`office_template/offer/upload` 返回的 `info.occurrences`
只是机械提取文档里所有 `{...}`，**不做任何有效性校验**——实测漏写 `[id]` 的自定义
字段和根本不存在的字段名都会照样计入 occurrences。所以「上传成功 + occurrences 有
条目」不代表变量能被替换，合法性只能在本地校验。

用法：
  # 校验 docx，字段清单实时从 current 租户拉取
  uv run --with python-docx python3 validate_placeholders.py --docx offer.docx

  # 字段清单用已保存的响应（离线校验 / 复现问题）
  python3 validate_placeholders.py --docx offer.docx --fields-json fields.json

  # 只校验一组占位符（不读 docx，无需 python-docx）
  python3 validate_placeholders.py --placeholders '候选人姓名,Annual%20Leave[111002]'

  # 电子签模板：允许签署区变量，并检查必需的个人签署区
  python3 validate_placeholders.py --docx offer.docx --template-type 1

输出（stdout）：JSON
  {
    "ok": true,
    "org_id": "...",
    "checked": 6,
    "invalid_count": 0,
    "placeholders": [
      {"raw": "候选人姓名", "kind": "standard", "valid": true, "note": null},
      {"raw": "Annual%20Leave[111002]", "kind": "custom", "field_id": 111002,
       "field_name": "Annual Leave", "valid": true, "note": null}
    ],
    "warnings": [...]
  }

失败时 exit code 1，且 ok 为 false。
"""

import argparse
import json
import os
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from placeholder_spec import (  # noqa: E402
    NON_PLACEHOLDER_FIELD_TYPES,
    STANDARD_SIGNATURE,
    STANDARD_UNVERIFIED,
    STANDARD_VARIABLES,
    decode_placeholder_name,
    encode_placeholder_name,
)

FIELDS_URL = "/api/outer/ats-offer/customFields/listCustomFieldsByOrgIdPermission"
ORG_INFO_URL = "/api/v2/org/info"

PLACEHOLDER_RE = re.compile(r"\{([^{}]*)\}")
# 占位符尾部的 [123]；ID 只允许纯数字
FIELD_ID_RE = re.compile(r"^(?P<name>.*)\[(?P<id>\d+)\]$")


def _require_no_profile_override() -> None:
    if os.environ.get("CLLMK_PROFILE", "").strip():
        raise SystemExit(
            "error: CLLMK_PROFILE is set; unset it, run 'cllmk auth switch', "
            "verify with 'cllmk auth status', then retry"
        )


def _cllmk_post(url: str, payload: str = "{}") -> dict:
    _require_no_profile_override()
    proc = subprocess.run(
        ["cllmk", "curl", "--url", url, "--method", "POST", "--payload", payload],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if proc.returncode != 0:
        raise SystemExit(f"error: {url} failed: {proc.stdout.strip()[:300]}")
    return json.loads(proc.stdout)


def fetch_fields() -> tuple[list[dict], str]:
    """拉取 current 租户的 Offer 字段清单，返回 (fields, orgId)。"""
    resp = _cllmk_post(FIELDS_URL)
    inner = resp.get("data") or {}
    if not inner.get("success", False):
        raise SystemExit(f"error: field list rejected: {json.dumps(inner, ensure_ascii=False)[:300]}")
    fields = inner.get("data") or []
    org_id = fields[0].get("orgId", "") if fields else ""
    return fields, org_id


def load_fields_json(path: str) -> tuple[list[dict], str]:
    """从已保存的响应读字段清单。兼容 cllmk 包装层与裸响应两种形态。"""
    raw = json.loads(Path(path).read_text(encoding="utf-8"))
    node = raw
    # cllmk 输出把服务端响应包在 data 里；服务端响应自身也有 data。逐层下钻找数组。
    for _ in range(4):
        if isinstance(node, list):
            break
        if not isinstance(node, dict):
            raise SystemExit(f"error: cannot locate a field array in {path}")
        node = node.get("data")
    if not isinstance(node, list):
        raise SystemExit(f"error: cannot locate a field array in {path}")
    org_id = node[0].get("orgId", "") if node else ""
    return node, org_id


def extract_from_docx(path: str) -> list[str]:
    """抽取 docx 正文、表格、页眉页脚里的占位符原文（去掉花括号）。"""
    try:
        from docx import Document
    except ImportError:
        raise SystemExit(
            "error: python-docx is required for --docx; rerun with "
            "'uv run --with python-docx python3 validate_placeholders.py ...'"
        )

    doc = Document(path)
    texts: list[str] = []

    def collect_container(container) -> None:
        for para in getattr(container, "paragraphs", []):
            texts.append(para.text)
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    collect_container(cell)

    collect_container(doc)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if part is not None:
                collect_container(part)

    found: list[str] = []
    for text in texts:
        found.extend(PLACEHOLDER_RE.findall(text))
    return found


def classify(raw: str, by_id: dict[int, dict], by_name: dict[str, list[dict]]) -> dict:
    """判定单个占位符。raw 是去掉花括号后的原文。"""
    result: dict = {"raw": raw, "kind": None, "valid": False, "note": None}

    match = FIELD_ID_RE.match(raw)
    if match:
        encoded_name = match.group("name")
        field_id = int(match.group("id"))
        decoded = decode_placeholder_name(encoded_name)
        result.update({"kind": "custom", "field_id": field_id, "field_name": decoded})

        field = by_id.get(field_id)
        if field is None:
            result["note"] = (
                f"字段 ID {field_id} 不在当前字段清单里。两种常见原因："
                "① 会话 hireMode 与占位符所属场景不符——Offer 字段按社招/校招分身，"
                "两边 ID 集合几乎不相交，先查 org/info 的 currentHireMode 确认；"
                "② 字段 ID 来自别的租户——字段 ID 是租户内标识，禁止跨租户复用。"
            )
            return result
        actual = field.get("name", "")
        if actual != decoded:
            result["note"] = (
                f"字段 ID {field_id} 的实际名称是 {actual!r}，占位符里写的是 {decoded!r}；"
                f"应写 {{{encode_placeholder_name(actual)}[{field_id}]}}"
            )
            return result
        if " " in decoded and "%20" not in encoded_name:
            result["note"] = (
                f"字段名含空格但未编码；应写 {{{encode_placeholder_name(actual)}[{field_id}]}}"
            )
            return result
        if field.get("type") in NON_PLACEHOLDER_FIELD_TYPES:
            result["note"] = (
                f"字段 type={field.get('type')} 不能作为占位符"
                "（附件 / 人员选择 / 招聘需求 / 附件发送开关类字段被 UI 排除）"
            )
            return result
        if field.get("isVisible") is False:
            result["valid"] = True
            result["note"] = "字段当前处于隐藏状态（isVisible=false），占位符合法但值可能为空"
            return result
        result["valid"] = True
        return result

    # 无 [id]：只能是标准变量
    decoded = decode_placeholder_name(raw)
    if decoded in STANDARD_VARIABLES:
        result.update({"kind": "standard", "valid": True})
        if decoded in STANDARD_UNVERIFIED:
            result["note"] = (
                "该变量在前端枚举表里，但不在 UI「查看字段」列表中且替换未经实测；"
                "使用前需向用户明示可用性未验证"
            )
        elif decoded in STANDARD_SIGNATURE:
            result["note"] = "电子签变量，仅在 templateType != 0 的模板里有意义"
        return result

    # 不是标准变量：大概率是漏了 [id] 的自定义字段
    candidates = by_name.get(decoded, [])
    if candidates:
        result["kind"] = "missing_field_id"
        suggestions = ", ".join(
            f"{{{encode_placeholder_name(c['name'])}[{c['id']}]}}" for c in candidates
        )
        plural = "个同名字段" if len(candidates) > 1 else "个字段"
        result["note"] = (
            f"{decoded!r} 不是标准变量，但匹配到 {len(candidates)} {plural}；"
            f"自定义字段必须带字段 ID：{suggestions}"
        )
        return result

    result["kind"] = "unknown"
    result["note"] = (
        f"{decoded!r} 既不是标准变量，也不匹配当前租户的任何 Offer 字段名。"
        "服务端不会拒绝它（occurrences 照样收录），但发 Offer 时不会被替换。"
    )
    return result


def main() -> int:
    parser = argparse.ArgumentParser(
        description="校验 Offer 附件模板占位符合法性（occurrences 不做校验，必须本地查）"
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--docx", help="待校验的 docx 路径")
    source.add_argument(
        "--placeholders",
        help="逗号分隔的占位符清单，不含花括号（如 '候选人姓名,Annual%%20Leave[111002]'）",
    )
    parser.add_argument(
        "--fields-json",
        help="已保存的字段清单响应；省略时实时调用 cllmk 拉取 current 租户",
    )
    parser.add_argument(
        "--template-type",
        type=int,
        default=0,
        help="模板的 templateType（0=不支持电子签）。非 0 时检查是否插入了个人签署区",
    )
    args = parser.parse_args()

    if args.fields_json:
        fields, org_id = load_fields_json(args.fields_json)
    else:
        fields, org_id = fetch_fields()

    by_id = {f["id"]: f for f in fields if isinstance(f.get("id"), int)}
    by_name: dict[str, list[dict]] = {}
    for f in fields:
        by_name.setdefault(f.get("name", ""), []).append(f)

    if args.docx:
        raws = extract_from_docx(args.docx)
    else:
        raws = [p.strip() for p in args.placeholders.split(",") if p.strip()]

    checked = [classify(raw, by_id, by_name) for raw in raws]
    invalid = [c for c in checked if not c["valid"]]

    warnings: list[str] = []
    if not raws:
        warnings.append(
            "没有提取到任何占位符。若文档里确实写了 {…}，最可能的原因是 Word/WPS 把占位符"
            "拆进了多个 run；用 build_offer_docx.py 生成可避免。"
        )
    names = {c["raw"] for c in checked}
    if args.template_type != 0 and "个人签署区" not in names:
        warnings.append(
            "templateType != 0（支持电子签）但文档里没有 {个人签署区}；"
            "候选人侧无签署位置。企业签署另需在「授权印章管理」配置印章。"
        )
    if args.template_type == 0:
        signature_used = [c["raw"] for c in checked if decode_placeholder_name(c["raw"]) in STANDARD_SIGNATURE]
        if signature_used:
            warnings.append(
                f"templateType=0（不支持电子签）但文档里含签署类变量 {signature_used}；"
                "这些变量不会生效。"
            )

    report = {
        "ok": not invalid,
        "org_id": org_id,
        "template_type": args.template_type,
        "checked": len(checked),
        "invalid_count": len(invalid),
        "placeholders": checked,
        "warnings": warnings,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if invalid else 0


if __name__ == "__main__":
    sys.exit(main())
