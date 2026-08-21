"""Tests for parse_form_file.py.

Run:
  uv run --with openpyxl --with python-docx --with pytest \
    pytest skills/ats/scripts/_test/test_parse_form_file.py -v
"""

from pathlib import Path
import sys

import pytest

# 被测脚本位于 _test/ 旁同级业务目录 `candidate/`
_PARSE_SCRIPT_DIR = Path(__file__).resolve().parent.parent / "candidate"
sys.path.insert(0, str(_PARSE_SCRIPT_DIR))

from parse_form_file import parse


# 测试样本在仓库根 scripts/test-samples 下
REPO_ROOT = Path(__file__).resolve().parents[4]
SAMPLES = REPO_ROOT / "scripts" / "test-samples"
DOCX_RESUME = SAMPLES / "标准简历&信息采集.docx"
XLSX_LIST = SAMPLES / "配置整理表 - CR Application form.xlsx"
XLSX_LAYOUT = SAMPLES / "新员工登记表.xlsx"


# ---------------------------------------------------------------------------
# docx: 标准简历&信息采集.docx
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def docx_result():
    return parse(str(DOCX_RESUME))


def test_docx_ok(docx_result):
    assert docx_result["ok"] is True


def test_docx_block_count(docx_result):
    # 1 (untitled 个人信息) + 4 (paragraphs: 工作经历/教育经历/家庭情况/附加信息)
    assert len(docx_result["blocks"]) == 5


def test_docx_first_block_default_title(docx_result):
    b0 = docx_result["blocks"][0]
    assert b0["title"] == "个人信息"


def test_docx_block_titles(docx_result):
    titles = [b["title"] for b in docx_result["blocks"]]
    assert titles == ["个人信息", "工作经历", "教育经历", "家庭情况", "附加信息"]


def _find_field(block, name):
    for f in block["fields"]:
        if f["name"] == name:
            return f
    return None


def test_docx_personal_basic_field_姓名(docx_result):
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "姓名")
    assert f is not None
    assert f["type_hint"] == "string"
    assert f["required"] is True


def test_docx_personal_性别_select(docx_result):
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "性别")
    assert f is not None
    assert f["type_hint"] == "select"


def test_docx_personal_出生日期_day(docx_result):
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "出生日期")
    assert f["type_hint"] == "day"


def test_docx_personal_期望月薪_number_bilingual(docx_result):
    """Bilingual field name split + number type from value '24000'."""
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "期望月薪")
    assert f is not None
    assert f["name_en"] == "Expected Monthly Salary"
    assert f["type_hint"] == "number"


def test_docx_personal_户口性质_bilingual_split(docx_result):
    """户口性质Domicile Type → name=户口性质, name_en=Domicile Type."""
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "户口性质")
    assert f is not None
    assert f["name_en"] == "Domicile Type"


def test_docx_personal_预计加盟日期_day(docx_result):
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "预计加盟日期")
    assert f is not None
    assert f["name_en"] == "When can join"
    assert f["type_hint"] == "day"


def test_docx_personal_有无子女_select_with_options(docx_result):
    """'下拉菜单\\n有Yes/无 No' → select with options."""
    b0 = docx_result["blocks"][0]
    f = _find_field(b0, "有无子女")
    assert f is not None
    assert f["name_en"] == "Children"
    assert f["type_hint"] == "select"


def test_docx_workexp_block_repeatable(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "工作经历")
    assert block["repeatable"] is True


def test_docx_workexp_field_薪资_bilingual_number(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "工作经历")
    f = _find_field(block, "薪资")
    assert f is not None
    assert f["name_en"] == "Salary"
    assert f["type_hint"] == "number"


def test_docx_workexp_named_fields_present(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "工作经历")
    names = {f["name"] for f in block["fields"]}
    assert {"薪资", "证明人职位", "证明人联系方式", "与证明人关系"} <= names


def test_docx_education_unnamed_lines_surfaced(docx_result):
    """教育经历 cell is freeform (no '：') — should surface as unnamed_lines."""
    block = next(b for b in docx_result["blocks"] if b["title"] == "教育经历")
    assert block["fields"] == []
    assert any("沈阳工业大学" in line for line in block["unnamed_lines"])


def test_docx_family_field_年龄_number(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "家庭情况")
    f = _find_field(block, "年龄")
    assert f is not None
    assert f["type_hint"] == "number"


def test_docx_family_field_当前所在城市_with_space_bilingual(docx_result):
    """'当前所在城市 Current city' (with space) — split bilingual."""
    block = next(b for b in docx_result["blocks"] if b["title"] == "家庭情况")
    f = _find_field(block, "当前所在城市")
    assert f is not None
    assert f["name_en"] == "Current city"


def test_docx_extra_yesno_question_is_bool(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "附加信息")
    f = _find_field(block, "您是否有亲戚或朋友在本公司任职")
    assert f is not None
    assert f["type_hint"] == "bool"


def test_docx_extra_followup_text_is_string(docx_result):
    block = next(b for b in docx_result["blocks"] if b["title"] == "附加信息")
    f = _find_field(block, "如有，请提供姓名与您的关系")
    assert f is not None
    assert f["type_hint"] == "string"


def test_docx_extra_declaration_is_confirm_not_day(docx_result):
    """『入职前声明』应识别为 confirm，不应因含『入职』被误判为 day."""
    block = next(b for b in docx_result["blocks"] if b["title"] == "附加信息")
    # bilingual split should drop trailing '/'
    f = _find_field(block, "入职前声明")
    assert f is not None, "入职前声明 字段应存在（双语拆分后无末尾 /）"
    assert f["name_en"] == "Pre-Employment Statements"
    assert f["type_hint"] == "confirm"


def test_docx_extra_yesno_about_agreement_is_bool(docx_result):
    """『你是否曾签署竞业限制协议』+ Yes/No 值 → bool；
    不应因字段名含『协议』被误判为 confirm."""
    block = next(b for b in docx_result["blocks"] if b["title"] == "附加信息")
    f = next((x for x in block["fields"] if "竞业限制协议" in x["name"]), None)
    assert f is not None
    assert f["type_hint"] == "bool"


# ---------------------------------------------------------------------------
# xlsx (clean list form): 配置整理表 - CR Application form.xlsx
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def xlsx_list_result():
    return parse(str(XLSX_LIST))


def test_xlsx_list_ok(xlsx_list_result):
    assert xlsx_list_result["ok"] is True


def test_xlsx_list_has_multiple_blocks(xlsx_list_result):
    # File has ≥ 3 Chapter rows in the main sheet
    assert len(xlsx_list_result["blocks"]) >= 3


def test_xlsx_list_first_block_title_bilingual(xlsx_list_result):
    b0 = xlsx_list_result["blocks"][0]
    assert b0["title"] == "个人信息与联系方式"
    assert b0["title_en"] == "Personal and contact details"


def _flatten_fields(result):
    return [f for b in result["blocks"] for f in b["fields"]]


def test_xlsx_list_field_中文名字(xlsx_list_result):
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if "中文名字" in x["name"]), None)
    assert f is not None
    assert f["name_en"].startswith("Given name in Chinese")
    assert f["type_hint"] == "string"
    assert f["required"] is True


def test_xlsx_list_field_性别_select(xlsx_list_result):
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if x["name"] == "性别 *" or x["name"] == "性别"), None)
    assert f is not None
    assert f["type_hint"] == "select"


def test_xlsx_list_field_生日_day(xlsx_list_result):
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if x["name"] == "生日 *" or x["name"] == "生日"), None)
    assert f is not None
    assert f["type_hint"] == "day"


def test_xlsx_list_field_累积绩点_number(xlsx_list_result):
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if "累积绩点" in x["name"]), None)
    assert f is not None
    assert f["type_hint"] == "number"


def test_xlsx_list_required_yes(xlsx_list_result):
    """H3 = YES → required True."""
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if "中文名字" in x["name"]), None)
    assert f["required"] is True


def test_xlsx_list_required_n(xlsx_list_result):
    """H33 = N → required False (若为其它，请注明)."""
    fields = _flatten_fields(xlsx_list_result)
    candidates = [x for x in fields if "若为其它" in x["name"] or "若其它" in x["name"]]
    assert candidates, "expected at least one 若为其它字段"
    assert any(x["required"] is False for x in candidates)


def test_xlsx_list_trigger_condition_extracted(xlsx_list_result):
    """Row with E25 = '当...为博士时' → trigger_condition non-null."""
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if x["name"] == "博士就读院校"), None)
    assert f is not None
    assert f["trigger_condition"] is not None
    assert "博士" in f["trigger_condition"]


def test_xlsx_list_trigger_condition_normalized_for_无(xlsx_list_result):
    """Row with E16='无' → trigger_condition normalized to None."""
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if "感兴趣的工作机会" in x["name"]), None)
    assert f is not None
    assert f["trigger_condition"] is None


def test_xlsx_list_picklist_reference_in_note(xlsx_list_result):
    """Row 16 G='See picklist' should be preserved as note."""
    fields = _flatten_fields(xlsx_list_result)
    f = next((x for x in fields if "感兴趣的工作机会" in x["name"]), None)
    assert f is not None
    note = (f.get("note") or "").lower()
    assert "picklist" in note or "see picklist" in note


# ---------------------------------------------------------------------------
# xlsx (layout form, fallback case): 新员工登记表.xlsx
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def xlsx_layout_result():
    return parse(str(XLSX_LAYOUT))


def test_xlsx_layout_falls_back(xlsx_layout_result):
    assert xlsx_layout_result["ok"] is False
    assert "fallback_text" in xlsx_layout_result
    assert xlsx_layout_result["reason"]


def test_xlsx_layout_fallback_contains_key_labels(xlsx_layout_result):
    txt = xlsx_layout_result["fallback_text"]
    # spot-check that several distinctive labels survived as text
    for label in ["员工个人信息记录表", "姓名", "身份证号码", "紧急联络人"]:
        assert label in txt, f"expected '{label}' in fallback_text"


# ---------------------------------------------------------------------------
# md passthrough
# ---------------------------------------------------------------------------


def test_md_input_returns_text(tmp_path):
    md = tmp_path / "sample.md"
    md.write_text("## 个人信息\n- 姓名 *\n- 手机号 *\n", encoding="utf-8")
    result = parse(str(md))
    assert result["ok"] is True
    assert result["format"] == "md"
    assert "姓名" in result["text"]


# ---------------------------------------------------------------------------
# unsupported / missing
# ---------------------------------------------------------------------------


def test_unsupported_extension(tmp_path):
    f = tmp_path / "sample.pdf"
    f.write_bytes(b"%PDF-fake")
    result = parse(str(f))
    assert result["ok"] is False
    assert "pdf" in result["reason"].lower() or "unsupported" in result["reason"].lower()


def test_missing_file():
    result = parse("/tmp/__no_such_file__.docx")
    assert result["ok"] is False
    assert "not found" in result["reason"].lower() or "missing" in result["reason"].lower()


# ---------------------------------------------------------------------------
# Edge cases / robustness
# ---------------------------------------------------------------------------


def test_empty_md(tmp_path):
    md = tmp_path / "empty.md"
    md.write_text("", encoding="utf-8")
    result = parse(str(md))
    assert result["ok"] is True
    assert result["text"] == ""


def test_empty_xlsx_falls_back_gracefully(tmp_path):
    """A blank xlsx with no recognizable header should fallback, not crash."""
    import openpyxl
    xlsx = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(str(xlsx))
    result = parse(str(xlsx))
    assert result["ok"] is False
    assert "fallback_text" in result


def test_empty_docx_no_tables(tmp_path):
    """A docx with only paragraphs (no tables) returns ok=True with empty blocks."""
    from docx import Document
    docx_path = tmp_path / "empty.docx"
    doc = Document()
    doc.add_paragraph("一份说明文档")
    doc.save(str(docx_path))
    result = parse(str(docx_path))
    assert result["ok"] is True
    assert result["blocks"] == []


# ---------------------------------------------------------------------------
# CLI end-to-end
# ---------------------------------------------------------------------------


def test_cli_runs_and_outputs_json(tmp_path):
    """Run the script as CLI and verify it prints valid JSON."""
    import json
    import subprocess

    md = tmp_path / "sample.md"
    md.write_text("- 姓名 *\n- 手机号 *\n", encoding="utf-8")

    script = _PARSE_SCRIPT_DIR / "parse_form_file.py"
    result = subprocess.run(
        [sys.executable, str(script), str(md)],
        capture_output=True, text=True, timeout=10,
    )
    assert result.returncode == 0, f"stderr={result.stderr}"
    data = json.loads(result.stdout)
    assert data["ok"] is True
    assert data["format"] == "md"
